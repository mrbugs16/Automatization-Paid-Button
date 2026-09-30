"""Lee la matriz de pruebas (Excel) y escribe una copia con los resultados obtenidos."""
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from config import settings

COLUMNAS = {
    "ID": "id",
    "Modulo": "modulo",
    "Caso de Prueba": "caso",
    "Tipo": "tipo",
    "Prioridad": "prioridad",
    "Precondiciones": "precondiciones",
    "Requiere 3DS": "requiere_3ds",
    "Pasos a Seguir": "pasos",
    "Datos de Prueba": "datos",
    "Resultado Esperado": "esperado",
    "Resultado Obtenido": "obtenido",
    "Estado": "estado",
    "Observaciones": "observaciones",
    "Ejecucion": "ejecucion",
}

# Estados del reporte -> valores permitidos en la columna 'Estado' del Excel
ESTADO_EXCEL = {"Error": "Bloqueado", "Manual": "Pendiente"}

COLORES_ESTADO = {
    "Aprobado": "C6EFCE",
    "Fallido": "FFC7CE",
    "Bloqueado": "F4B084",
    "Pendiente": "FFEB9C",
    "Manual": "DDEBF7",
}


def cargar_matriz(ruta=settings.MATRIZ_XLSX):
    """Regresa {ID: {campo: valor}} en el orden de la matriz."""
    ws = load_workbook(ruta, read_only=True, data_only=True)[settings.HOJA_MATRIZ]
    filas = ws.iter_rows(values_only=True)
    encabezados = [COLUMNAS.get(str(h).strip(), str(h)) if h else None for h in next(filas)]
    casos = {}
    for fila in filas:
        registro = {k: (v if v is not None else "") for k, v in zip(encabezados, fila) if k}
        # Las filas con solo el ID son casos eliminados (no se renumeran): se ignoran
        if str(registro.get("id", "")).startswith("TC-") and registro.get("caso"):
            casos[registro["id"]] = registro
    return casos


def escribir_resultados(resultados, destino, origen=settings.MATRIZ_XLSX):
    """Copia la matriz a 'destino' llenando 'Resultado Obtenido' y 'Estado'."""
    wb = load_workbook(origen)
    ws = wb[settings.HOJA_MATRIZ]
    encabezados = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    col_id = encabezados["ID"]
    col_obtenido = encabezados["Resultado Obtenido"]
    col_estado = encabezados["Estado"]
    for fila in range(2, ws.max_row + 1):
        caso_id = ws.cell(fila, col_id).value
        r = resultados.get(caso_id)
        if not r:
            continue
        estado = ESTADO_EXCEL.get(r["estado"], r["estado"])
        prefijo = "[Prueba manual] " if r["estado"] == "Manual" else ""
        ws.cell(fila, col_obtenido).value = prefijo + r["obtenido"]
        celda = ws.cell(fila, col_estado)
        celda.value = estado
        if estado in COLORES_ESTADO:
            color = COLORES_ESTADO[estado]
            celda.fill = PatternFill("solid", start_color=color, end_color=color)
    wb.save(destino)
